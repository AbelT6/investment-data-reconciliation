"""
export_excel.py
Writes output/exception_log.xlsx with three sheets:
  Exceptions       - the exception log as an Excel Table (tblExceptions)
  Security Master  - reference data as an Excel Table (tblMaster), for XLOOKUP
  Summary          - counts and metrics from output/summary.csv

The analysis layer (pivot, XLOOKUP, conditional formatting, data validation)
is built by hand in Excel -- see README "Excel layer".
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
OUTPUT = ROOT / "output"

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0.00;[Red]-#,##0.00'


def write_sheet(ws, df: pd.DataFrame, table_name: str | None, money_cols=()) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])

    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
        ws.column_dimensions[letter].width = min(width, 60)
        if col in money_cols:
            for cell in ws[letter][1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = MONEY

    if table_name:
        ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
        t = Table(displayName=table_name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)


def main() -> None:
    log = pd.read_csv(OUTPUT / "exception_log.csv", dtype={"cusip": str})
    master = pd.read_csv(DATA / "security_master.csv", dtype={"cusip": str})
    summary = pd.read_csv(OUTPUT / "summary.csv")

    # Numeric-looking values in these columns stay numbers so Excel can format them;
    # CUSIPs stay text so Excel can't strip their zeros (the bug this project is about).
    numeric_rows = log.field.isin(["market_value", "book_value", "accrued_interest", "row_count"])
    for col in ["custodian_value", "accounting_value"]:
        log[col] = log[col].astype(object)
        log.loc[numeric_rows, col] = pd.to_numeric(log.loc[numeric_rows, col])

    wb = Workbook()
    write_sheet(wb.active, log, "tblExceptions", money_cols=("custodian_value", "accounting_value", "difference"))
    wb.active.title = "Exceptions"
    write_sheet(wb.create_sheet("Security Master"), master, "tblMaster")
    write_sheet(wb.create_sheet("Summary"), summary, None)

    # Force CUSIP cells to text format so re-saving in Excel can't strip leading zeros
    for ws in (wb["Exceptions"], wb["Security Master"]):
        col = [c.value for c in ws[1]].index("cusip") + 1
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = "@"

    out = OUTPUT / "exception_log.xlsx"
    wb.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
